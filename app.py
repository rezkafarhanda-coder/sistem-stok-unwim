import streamlit as st
import pandas as pd
import datetime
import base64
import io
import re
from pathlib import Path

# Import khusus untuk styling file Word
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

# Import khusus untuk Google Sheets
from streamlit_gsheets import GSheetsConnection

# ==================================================
# KONFIGURASI HALAMAN
# ==================================================
st.set_page_config(
    page_title="Sistem Stok Barang",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ==================================================
# CSS MODERN DASHBOARD
# ==================================================
st.markdown("""
<style>
.stApp{ background-color:#f4f6f9; }
.custom-header {
    background-color: #12715b; color: white; padding: 12px 25px; border-radius: 8px; margin-bottom: 25px; 
    display: flex; justify-content: space-between; align-items: center; box-shadow: 0 4px 10px rgba(0,0,0,0.1);
}
.header-left { display: flex; align-items: center; gap: 15px; }
.header-logo { height: 55px; width: auto; }
.header-title { display: flex; flex-direction: column; justify-content: center; }
.title-main { font-size: 22px; font-weight: bold; letter-spacing: 0.5px; margin: 0; color: white !important; }
.title-sub { font-size: 13px; font-weight: 400; letter-spacing: 0.5px; margin-top: 3px; color: rgba(255, 255, 255, 0.8) !important; }
.header-right { display: flex; align-items: center; gap: 20px; font-size: 14px; font-weight: 500; }
.header-item { display: flex; align-items: center; gap: 8px; }
.icon-svg { width: 18px; height: 18px; fill: none; stroke: currentColor; stroke-width: 2; stroke-linecap: round; stroke-linejoin: round; }

button[data-baseweb="tab"] p, button[data-baseweb="tab"] span { color: #111827 !important; font-weight: 600 !important; font-size: 16px !important; }
button[data-baseweb="tab"][aria-selected="true"] p, button[data-baseweb="tab"][aria-selected="true"] span { font-weight: 800 !important; }
div[data-baseweb="tab-highlight"] { background-color: #12715b !important; }

.card-ringkasan-baru { border-radius: 14px; padding: 18px 20px; margin-bottom: 15px; display: flex; align-items: flex-start; gap: 16px; position: relative; overflow: hidden; box-shadow: 0 2px 8px rgba(0,0,0,0.05); }
.card-ringkasan-baru::before { content: ""; position: absolute; width: 140px; height: 140px; bottom: -50px; left: -50px; border-radius: 45%; z-index: 0; pointer-events: none; }
.card-ringkasan-baru::after { content: ""; position: absolute; width: 160px; height: 160px; bottom: -70px; left: -40px; border-radius: 40%; z-index: 0; pointer-events: none; opacity: 0.7; }
.card-ringkasan-baru.hijau { background: linear-gradient(135deg, #f2faf7 0%, #e3f5ee 100%); }
.card-ringkasan-baru.hijau::before { background: rgba(32, 201, 151, 0.18); transform: rotate(15deg); }
.card-ringkasan-baru.hijau::after { background: rgba(18, 113, 91, 0.08); transform: rotate(45deg); }
.card-ringkasan-baru.kuning { background: linear-gradient(135deg, #fff9ee 0%, #fff2d9 100%); }
.card-ringkasan-baru.kuning::before { background: rgba(255, 176, 32, 0.16); transform: rotate(25deg); }
.card-ringkasan-baru.kuning::after { background: rgba(255, 143, 0, 0.07); transform: rotate(55deg); }

.circle-icon-baru, .content-kanan-baru { position: relative; z-index: 1; }
.circle-icon-baru { width: 54px; height: 54px; border-radius: 50%; display: flex; justify-content: center; align-items: center; font-size: 24px; color: white; flex-shrink: 0; }
.circle-icon-baru.hijau { background-color: #12715b; }
.circle-icon-baru.kuning { background-color: #ffb020; }
.content-kanan-baru { display: flex; flex-direction: column; }
.title-card-baru { font-size: 13px; font-weight: bold; letter-spacing: 0.5px; margin: 0; }
.angka-card-baru { font-size: 36px; font-weight: bold; line-height: 1.1; margin: 3px 0; }
.desc-card-baru { font-size: 12px; color: #4b5563; line-height: 1.4; margin: 0; }

div[data-baseweb="input"]{ border:1px solid #cbd5e1 !important; border-radius:8px !important; background:white !important; }
div[data-baseweb="select"] > div{ border:1px solid #cbd5e1 !important; border-radius:8px !important; background:white !important; }

.table-gudang { width: 100%; border-collapse: separate !important; border-spacing: 0 !important; background: white; border-radius: 12px !important; overflow: hidden; border: 1px solid #e5e7eb !important; }
.table-gudang th { background: #12715b !important; color: white !important; padding: 14px 12px !important; text-align: center !important; font-weight: bold !important; font-size: 14px !important; border: 1px solid #106652 !important; }
.table-gudang td { padding: 12px !important; text-align: center !important; vertical-align: middle !important; border: 1px solid #eef2f5 !important; color: #374151 !important; font-size: 14px !important; }
.table-gudang tr:nth-child(even) { background: #ffffff; }

.status-badge-container { display: flex; justify-content: center; align-items: center; }
.status-badge { display: inline-flex; align-items: center; gap: 8px; padding: 6px 16px; border-radius: 6px; font-size: 13px; font-weight: 500; min-width: 135px; }
.status-badge.good { background-color: #eefaf6; color: #12715b; border: 1px solid #d1f2e5; }
.status-badge.good .badge-dot { background-color: #22c55e; }
.status-badge.low { background-color: #fff9ed; color: #ffb020; border: 1px solid #ffecc7; }
.status-badge.low .badge-dot { background-color: #ffb020; }
.status-badge.out { background-color: #fef2f2; color: #ef4444; border: 1px solid #fee2e2; }
.status-badge.out .badge-dot { background-color: #ef4444; }
.badge-dot { width: 8px; height: 8px; border-radius: 50%; display: inline-block; flex-shrink: 0; }

div.stButton > button{ border-radius:10px !important; font-weight:bold !important; }
div.element-container:has(.marker-hijau) + div.element-container button{ background:#22c55e !important; color:white !important; border: none !important;}
div.element-container:has(.marker-merah) + div.element-container button{ background:#ef4444 !important; color:white !important; border: none !important;}
div.element-container:has(.marker-teal) + div.element-container { margin-top: 12px !important;  }
div.element-container:has(.marker-teal) + div.element-container button { background: #12715b !important; color: white !important; border: none !important; border-radius: 8px !important; height: 40px !important; display: flex !important; align-items: center !important; justify-content: center !important; }
div.element-container:has(.marker-teal) + div.element-container button:hover { background: #0f766e !important; color: white !important; }

/* Styling khusus untuk form container */
div[data-testid="stForm"] { border: 1px solid #e5e7eb !important; border-radius: 12px !important; padding: 20px !important; background-color: #ffffff !important; box-shadow: 0 2px 5px rgba(0,0,0,0.02); }
</style>
""", unsafe_allow_html=True)

# ==================================================
# LOGIKA PROSES LOGO
# ==================================================
def get_image_as_base64(image_path):
    try:
        if not Path(image_path).is_file(): return ""
        with open(image_path, "rb") as image_file:
            return f"data:image/png;base64,{base64.b64encode(image_file.read()).decode('utf-8')}"
    except Exception: return ""

nama_file_logo = r"C:\Users\Mochammad Rezka\Pictures\Logo Unwim.png"
logo_base64 = get_image_as_base64(nama_file_logo)
display_style = "display: block;" if logo_base64 else "display: none;"

sekarang = datetime.datetime.now()
bulan_indo = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
tanggal_otomatis = f"{sekarang.day:02d} {bulan_indo[sekarang.month-1]} {sekarang.year}"

# ==================================================
# HEADER TAMPILAN
# ==================================================
st.markdown(f"""
<div class="custom-header">
<div class="header-left">
<img src="{logo_base64}" class="header-logo" alt="Logo" style="{display_style}">
<div class="header-title">
<div class="title-main">SISTEM STOK BARANG</div>
<div class="title-sub">REKTORAT UNIVERSITAS WINAYA MUKTI</div>
</div>
</div>
<div class="header-right">
<div class="header-item">
<svg class="icon-svg" viewBox="0 0 24 24">
<rect x="3" y="4" width="18" height="18" rx="2" ry="2"></rect>
<line x1="16" y1="2" x2="16" y2="6"></line>
<line x1="8" y1="2" x2="8" y2="6"></line>
<line x1="3" y1="10" x2="21" y2="10"></line>
</svg>
{tanggal_otomatis}
</div>
</div>
</div>
""", unsafe_allow_html=True)

# ==================================================
# KONEKSI DATABASE (GOOGLE SHEETS)
# ==================================================
conn = st.connection("gsheets", type=GSheetsConnection)

if "data_ditarik" not in st.session_state:
    try:
        # Menarik data permanen dari Sheet
        st.session_state.df_stok = conn.read(worksheet="Stok", ttl=0).dropna(how="all")
        st.session_state.df_transaksi = conn.read(worksheet="Transaksi", ttl=0).dropna(how="all")
        
        # Konversi tipe data agar tidak error saat dihitung
        st.session_state.df_stok["ID Barang"] = st.session_state.df_stok["ID Barang"].astype(str)
        st.session_state.df_stok["Jumlah Stok"] = pd.to_numeric(st.session_state.df_stok["Jumlah Stok"], errors='coerce').fillna(0)
        
        st.session_state.data_ditarik = True
    except Exception as e:
        # Jika gagal (aplikasi belum disambungkan ke Sheet), gunakan data sampel
        st.session_state.df_stok = pd.DataFrame({
            "ID Barang":["1","2","3","4","5"],
            "Nama Barang":["Kertas A4", "Pena Pilot", "Tinta Printer", "Buku Tulis", "Spidol"],
            "Kategori":["ATK", "ATK", "Elektronik", "ATK", "ATK"],
            "Jumlah Stok":[35,15,5,20,0],
            "Satuan":["Rim", "Pcs", "Pcs", "Pcs", "Pcs"]
        })
        st.session_state.df_transaksi = pd.DataFrame(columns=[
            "Waktu", "Jenis", "ID Barang", "Nama Barang", "Jml Transaksi", "Pengambil"
        ])
        st.warning("⚠️ Aplikasi berjalan dalam mode lokal. Hubungkan ke Google Sheets untuk menyimpan data permanen.")

def status_stok(jumlah):
    if jumlah > 20: return '<div class="status-badge-container"><div class="status-badge good"><span class="badge-dot"></span>Good Stock</div></div>'
    elif jumlah > 0: return '<div class="status-badge-container"><div class="status-badge low"><span class="badge-dot"></span>Low Stock</div></div>'
    else: return '<div class="status-badge-container"><div class="status-badge out"><span class="badge-dot"></span>Out Of Stock</div></div>'

st.session_state.df_stok["Status"] = st.session_state.df_stok["Jumlah Stok"].apply(status_stok)

# Fungsi bantuan untuk menyimpan ke Google Sheets
def simpan_stok():
    try: conn.update(worksheet="Stok", data=st.session_state.df_stok)
    except: pass

def simpan_transaksi():
    try: conn.update(worksheet="Transaksi", data=st.session_state.df_transaksi)
    except: pass

# ==================================================
# PEMBUATAN TAB (SLIDE)
# ==================================================
tab1, tab2 = st.tabs(["📊 Dashboard Utama", "🕒 Riwayat Transaksi"])

# ==================================================
# SLIDE 1: DASHBOARD UTAMA
# ==================================================
with tab1:
    
    # 🌟 Notifikasi Tab 1
    if "notif_tab1" in st.session_state:
        st.success(st.session_state.notif_tab1)
        del st.session_state.notif_tab1

    col_kiri, col_tengah, col_kanan = st.columns([1.1, 3, 1.2], gap="large")

    with col_kiri:
        st.markdown('<h3 style="color: #12715b; margin: 10px 0 20px 0; font-size: 19px; font-weight: bold; letter-spacing: 0.5px;">RINGKASAN STOK</h3>', unsafe_allow_html=True)
        total_produk = len(st.session_state.df_stok)
        barang_langka = len(st.session_state.df_stok[st.session_state.df_stok["Jumlah Stok"] < 10])

        st.markdown(f"""
        <div class="card-ringkasan-baru hijau">
            <div class="circle-icon-baru hijau">📦</div>
            <div class="content-kanan-baru">
                <div class="title-card-baru" style="color: #12715b;">TOTAL MACAM PRODUK</div>
                <div class="angka-card-baru" style="color: #12715b;">{total_produk}</div>
            </div>
        </div>
        <div class="card-ringkasan-baru kuning">
            <div class="circle-icon-baru kuning">⚠️</div>
            <div class="content-kanan-baru">
                <div class="title-card-baru" style="color: #ff8f00;">BARANG HAMPIR HABIS</div>
                <div class="angka-card-baru" style="color: #ff8f00;">{barang_langka}</div>
                <div class="desc-card-baru">Segera Restock</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("### Grafik Stok")
        chart_data = st.session_state.df_stok[["Nama Barang","Jumlah Stok"]].set_index("Nama Barang")
        st.bar_chart(chart_data)

    with col_tengah:
        st.markdown("""
        <div style="display: flex; align-items: center; gap: 10px; margin-top: 10px; margin-bottom: 20px;">
            <svg width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="#12715b" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round">
                <path d="M21 16V8a2 2 0 0 0-1-1.73l-7-4a2 2 0 0 0-2 0l-7 4A2 2 0 0 0 3 8v8a2 2 0 0 0 1 1.73l7 4a2 2 0 0 0 2 0l7-4A2 2 0 0 0 21 16z"></path>
                <polyline points="3.27 6.96 12 12.01 20.73 6.96"></polyline>
                <line x1="12" y1="22.08" x2="12" y2="12"></line>
            </svg>
            <h3 style="color: #12715b; margin: 0; font-size: 20px; font-weight: bold; letter-spacing: 0.5px;">MANAJEMEN STOK</h3>
        </div>
        """, unsafe_allow_html=True)

        c1,c2,c3,c4,c5 = st.columns([2.0, 1.5, 0.8, 1.2, 1.5])
        with c1: input_nama = st.text_input("Nama Barang", placeholder="Ketik Disini", key="add_nama")
        with c2: input_kategori = st.text_input("Kategori", placeholder="Contoh: ATK", key="add_kat")
        with c3: input_qty = st.number_input("Jml", min_value=0, step=1, key="add_qty")
        with c4: input_satuan = st.selectbox("Satuan", ["Pcs","Rim","Lusin"], key="add_satuan")
        with c5:
            st.markdown('<div class="marker-teal" style="display:none;"></div>', unsafe_allow_html=True)
            btn_add = st.button("+ Add Item", use_container_width=True, key="btn_add")

        if btn_add and input_nama != "":
            id_baru = str(len(st.session_state.df_stok) + 1)
            data_baru = pd.DataFrame({
                "ID Barang":[id_baru], "Nama Barang":[input_nama],
                "Kategori":[input_kategori], "Jumlah Stok":[input_qty],
                "Satuan":[input_satuan], "Status":[""]
            })
            st.session_state.df_stok = pd.concat([st.session_state.df_stok, data_baru], ignore_index=True)
            
            simpan_stok() # Simpan ke Google Sheets
            
            st.session_state.notif_tab1 = f"✅ Barang '{input_nama}' berhasil ditambahkan ke gudang."
            st.rerun()

        # ==================================================
        # EDIT MASTER BARANG (PILIH NAMA BARANG)
        # ==================================================
        is_editing_master = st.session_state.get("edit_master_nama", "-") != "-"
        
        with st.expander("✏️ Edit / Hapus Master Barang", expanded=is_editing_master):
            pilihan_barang = ["-"] + list(st.session_state.df_stok["Nama Barang"])
            edit_nama_pilih = st.selectbox("Pilih Nama Barang:", pilihan_barang, key="edit_master_nama")
            
            if edit_nama_pilih != "-":
                idx_edit = st.session_state.df_stok.index[st.session_state.df_stok["Nama Barang"] == edit_nama_pilih].tolist()[0]
                data_edit = st.session_state.df_stok.iloc[idx_edit]
                
                e1, e2, e3, e4 = st.columns([2.0, 1.5, 0.8, 1.2])
                with e1: edit_nama = st.text_input("Nama Baru", data_edit["Nama Barang"], key="e_nama")
                with e2: edit_kat = st.text_input("Kategori Baru", data_edit["Kategori"], key="e_kat")
                with e3: edit_qty = st.number_input("Jml Baru", value=int(data_edit["Jumlah Stok"]), key="e_qty")
                with e4: 
                    try: idx_satuan = ["Pcs","Rim","Lusin"].index(data_edit["Satuan"])
                    except ValueError: idx_satuan = 0
                    edit_satuan = st.selectbox("Satuan Baru", ["Pcs","Rim","Lusin"], index=idx_satuan, key="e_satuan")
                
                st.markdown("<br>", unsafe_allow_html=True)
                ed_b1, ed_b2 = st.columns(2)
                
                with ed_b1:
                    st.markdown('<div class="marker-hijau" style="display:none;"></div>', unsafe_allow_html=True)
                    if st.button("💾 Simpan Perubahan", use_container_width=True, key="btn_simpan_master"):
                        # Hanya perbarui data master barang
                        st.session_state.df_stok.at[idx_edit, "Nama Barang"] = edit_nama
                        st.session_state.df_stok.at[idx_edit, "Kategori"] = edit_kat
                        st.session_state.df_stok.at[idx_edit, "Jumlah Stok"] = edit_qty
                        st.session_state.df_stok.at[idx_edit, "Satuan"] = edit_satuan
                        
                        simpan_stok() # Simpan ke Google Sheets
                        
                        st.session_state.notif_tab1 = f"✅ Data '{edit_nama}' berhasil diubah."
                        if "edit_master_nama" in st.session_state:
                            del st.session_state["edit_master_nama"]
                        st.rerun()
                        
                with ed_b2:
                    st.markdown('<div class="marker-merah" style="display:none;"></div>', unsafe_allow_html=True)
                    if st.button("🗑️ Hapus Barang", use_container_width=True, key="btn_hapus_master"):
                        nama_terhapus = data_edit["Nama Barang"]
                        st.session_state.df_stok = st.session_state.df_stok.drop(idx_edit).reset_index(drop=True)
                        
                        simpan_stok() # Simpan ke Google Sheets
                        
                        st.session_state.notif_tab1 = f"🗑️ Barang '{nama_terhapus}' berhasil dihapus dari sistem."
                        if "edit_master_nama" in st.session_state:
                            del st.session_state["edit_master_nama"]
                        st.rerun()

        st.markdown("<br>", unsafe_allow_html=True)
        tabel_html = st.session_state.df_stok.to_html(classes="table-gudang", index=False, escape=False)
        st.markdown(tabel_html, unsafe_allow_html=True)

    with col_kanan:
        st.markdown("""
        <div style="display: flex; align-items: center; gap: 8px; margin-bottom: 15px;">
            <svg width="22" height="22" viewBox="0 0 24 24" fill="none" stroke="#12715b" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round">
                <path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"></path>
                <polyline points="14 2 14 8 20 8"></polyline>
                <line x1="16" y1="13" x2="8" y2="13"></line>
                <line x1="16" y1="17" x2="8" y2="17"></line>
                <polyline points="10 9 9 9 8 9"></polyline>
            </svg>
            <h3 style="color: #12715b; margin: 0; font-size: 18px; font-weight: bold; letter-spacing: 0.5px;">TRANSAKSI</h3>
        </div>
        """, unsafe_allow_html=True)

        with st.form("form_transaksi", clear_on_submit=True):
            input_barcode = st.text_input("Scan Barcode / Input ID", placeholder="Ketik angka ID Barang (1-5)")
            qty = st.number_input("Jumlah", min_value=1, step=1)
            
            input_pengambil = st.selectbox(
                "Pengambil", ["-", "Yayasan", "Rektorat", "PMB", "FEB", "FAHUTAN", "FAPERTA", "FTPA"],
                help="Pilih fakultas pengambil untuk pencatatan riwayat transaksi."
            )

            st.markdown('<div class="marker-hijau" style="display:none;"></div>', unsafe_allow_html=True)
            btn_masuk = st.form_submit_button("📥 Barang Masuk", use_container_width=True)

            st.markdown('<div class="marker-merah" style="display:none;"></div>', unsafe_allow_html=True)
            btn_keluar = st.form_submit_button("📤 Barang Keluar", use_container_width=True)


        # Logika Barang Masuk
        if btn_masuk:
            if input_barcode in st.session_state.df_stok["ID Barang"].values:
                stok_sekarang = st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Jumlah Stok"].values[0]
                nama_brg = st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Nama Barang"].values[0]
                waktu_skrg = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Jumlah Stok"] += qty
                stok_terbaru = stok_sekarang + qty

                log_baru = pd.DataFrame({
                    "Waktu": [waktu_skrg],
                    "Jenis": ['<span style="background-color:#dcfce7; color:#166534; padding:4px 8px; border-radius:4px; font-weight:bold; font-size:12px;">Masuk</span>'],
                    "ID Barang": [input_barcode], "Nama Barang": [nama_brg],
                    "Jml Transaksi": [f"+ {qty}"], "Pengambil": [input_pengambil],
                    "Barang Tersedia": [stok_terbaru]
                })
                st.session_state.df_transaksi = pd.concat([log_baru, st.session_state.df_transaksi], ignore_index=True)
                
                simpan_stok()
                simpan_transaksi()
                
                st.session_state.notif_tab1 = f"✅ Sukses: {qty} {nama_brg} ditambahkan."
                st.rerun()
            else:
                st.error("❌ Gagal: ID Barang tidak ditemukan.")

        # Logika Barang Keluar
        if btn_keluar:
            if input_barcode in st.session_state.df_stok["ID Barang"].values:
                stok_sekarang = st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Jumlah Stok"].values[0]
                if stok_sekarang >= qty:
                    nama_brg = st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Nama Barang"].values[0]
                    waktu_skrg = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                    st.session_state.df_stok.loc[st.session_state.df_stok["ID Barang"] == input_barcode, "Jumlah Stok"] -= qty
                    stok_terbaru = stok_sekarang - qty

                    log_baru = pd.DataFrame({
                        "Waktu": [waktu_skrg],
                        "Jenis": ['<span style="background-color:#fee2e2; color:#991b1b; padding:4px 8px; border-radius:4px; font-weight:bold; font-size:12px;">Keluar</span>'],
                        "ID Barang": [input_barcode], "Nama Barang": [nama_brg],
                        "Jml Transaksi": [f"- {qty}"], "Pengambil": [input_pengambil],
                        "Barang Tersedia": [stok_terbaru]
                    })
                    st.session_state.df_transaksi = pd.concat([log_baru, st.session_state.df_transaksi], ignore_index=True)
                    
                    simpan_stok()
                    simpan_transaksi()
                    
                    st.session_state.notif_tab1 = f"✅ Sukses: {qty} {nama_brg} dikeluarkan."
                    st.rerun()
                else:
                    st.error(f"❌ Gagal: Stok {nama_brg} tidak mencukupi. (Sisa: {stok_sekarang})")
            else:
                st.error("❌ Gagal: ID Barang tidak ditemukan.")

# ==================================================
# SLIDE 2: RIWAYAT TRANSAKSI
# ==================================================
with tab2:
    
    if "notif_tab2" in st.session_state:
        st.success(st.session_state.notif_tab2)
        del st.session_state.notif_tab2

    st.markdown("""
    <div style="display: flex; align-items: center; gap: 10px; margin-top: 15px; margin-bottom: 15px;">
        <svg width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="#12715b" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round">
            <path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"></path>
            <polyline points="14 2 14 8 20 8"></polyline>
            <line x1="16" y1="13" x2="8" y2="13"></line>
            <line x1="16" y1="17" x2="8" y2="17"></line>
            <polyline points="10 9 9 9 8 9"></polyline>
        </svg>
        <h3 style="color: #12715b; margin: 0; font-size: 20px; font-weight: bold; letter-spacing: 0.5px;">RIWAYAT TRANSAKSI</h3>
    </div>
    """, unsafe_allow_html=True)

    f1, f2 = st.columns(2)
    with f1:
        rentang_tanggal = st.date_input("📅 Pilih Rentang Tanggal", value=(datetime.date.today(), datetime.date.today()), key="unik_filter_tanggal")
        if len(rentang_tanggal) == 2: tgl_mulai, tgl_akhir = rentang_tanggal
        else: tgl_mulai = tgl_akhir = rentang_tanggal[0]
            
    with f2:
        list_barang = ["Semua Barang"] + sorted(list(st.session_state.df_stok["Nama Barang"].unique()))
        filter_barang = st.selectbox("📦 Filter Nama Barang", list_barang, key="unik_filter_barang")

    if st.session_state.df_transaksi.empty:
        st.info("Belum ada transaksi yang tercatat di sistem. Silakan lakukan transaksi terlebih dahulu.")
    else:
        is_editing_tx = st.session_state.get("sel_edit_nama_tx", "-") != "-"
        
        with st.expander("✏️ Edit / Hapus Riwayat Transaksi", expanded=is_editing_tx):
            pilihan_barang_tx = ["-"] + list(st.session_state.df_transaksi["Nama Barang"].unique())
            edit_nama_tx = st.selectbox("1. Pilih Nama Barang:", pilihan_barang_tx, key="sel_edit_nama_tx")
            
            if edit_nama_tx != "-":
                df_tx_filtered = st.session_state.df_transaksi[st.session_state.df_transaksi["Nama Barang"] == edit_nama_tx]
                
                tx_options = ["-"]
                for i, r in df_tx_filtered.iterrows():
                    jenis_clean = re.sub(r'<[^>]+>', '', str(r['Jenis']))
                    tx_options.append(f"{i} | {r['Waktu']} | {jenis_clean} ({r['Jml Transaksi']})")
                    
                sel_tx = st.selectbox("2. Pilih Waktu Transaksi yang akan diedit:", tx_options, key="sel_edit_tx_detail")
                
                if sel_tx != "-":
                    idx_tx = int(sel_tx.split(" | ")[0])
                    row_tx = st.session_state.df_transaksi.loc[idx_tx]
                    
                    old_jenis_bersih = re.sub(r'<[^>]+>', '', str(row_tx['Jenis']))
                    old_qty = int(re.sub(r'\D', '', str(row_tx['Jml Transaksi'])))
                    
                    te1, te2, te3 = st.columns([1.5, 1.5, 1.5])
                    with te1:
                        idx_j = 0 if old_jenis_bersih == "Masuk" else 1
                        new_jenis = st.selectbox("Ubah Jenis", ["Masuk", "Keluar"], index=idx_j, key="e_tx_jenis")
                    with te2:
                        new_qty = st.number_input("Ubah Jumlah", value=old_qty, min_value=1, step=1, key="e_tx_qty")
                    with te3:
                        fakultas = ["-", "Yayasan", "Rektorat", "PMB", "FEB", "FAHUTAN", "FAPERTA", "FTPA"]
                        try: idx_p = fakultas.index(row_tx["Pengambil"])
                        except ValueError: idx_p = 0
                        new_pengambil = st.selectbox("Ubah Pengambil", fakultas, index=idx_p, key="e_tx_pengambil")
                    
                    st.markdown("<br>", unsafe_allow_html=True)
                    ed_tx1, ed_tx2 = st.columns(2)
                    
                    with ed_tx1:
                        st.markdown('<div class="marker-hijau" style="display:none;"></div>', unsafe_allow_html=True)
                        if st.button("💾 Simpan Perubahan Transaksi", use_container_width=True, key="btn_save_tx"):
                            id_brg = row_tx['ID Barang']
                            
                            if old_jenis_bersih == "Masuk":
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] -= old_qty
                            else:
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] += old_qty
                                
                            if new_jenis == "Masuk":
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] += new_qty
                                jenis_html = '<span style="background-color:#dcfce7; color:#166534; padding:4px 8px; border-radius:4px; font-weight:bold; font-size:12px;">Masuk</span>'
                                jml_str = f"+ {new_qty}"
                            else:
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] -= new_qty
                                jenis_html = '<span style="background-color:#fee2e2; color:#991b1b; padding:4px 8px; border-radius:4px; font-weight:bold; font-size:12px;">Keluar</span>'
                                jml_str = f"- {new_qty}"
                                
                            # Ambil stok terbaru setelah diedit
                            stok_terbaru_edit = st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'].values[0]

                            st.session_state.df_transaksi.at[idx_tx, 'Jenis'] = jenis_html
                            st.session_state.df_transaksi.at[idx_tx, 'Jml Transaksi'] = jml_str
                            st.session_state.df_transaksi.at[idx_tx, 'Pengambil'] = new_pengambil
                            st.session_state.df_transaksi.at[idx_tx, 'Barang Tersedia'] = stok_terbaru_edit
                            
                            simpan_stok()
                            simpan_transaksi()
                            
                            st.session_state.notif_tab2 = "✅ Riwayat Transaksi berhasil diubah dan Stok Master telah disesuaikan."
                            if "sel_edit_nama_tx" in st.session_state: del st.session_state["sel_edit_nama_tx"]
                            st.rerun()

                    with ed_tx2:
                        st.markdown('<div class="marker-merah" style="display:none;"></div>', unsafe_allow_html=True)
                        if st.button("🗑️ Hapus Transaksi", use_container_width=True, key="btn_del_tx"):
                            id_brg = row_tx['ID Barang']
                            
                            if old_jenis_bersih == "Masuk":
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] -= old_qty
                            else:
                                st.session_state.df_stok.loc[st.session_state.df_stok['ID Barang'] == id_brg, 'Jumlah Stok'] += old_qty
                                
                            st.session_state.df_transaksi = st.session_state.df_transaksi.drop(idx_tx).reset_index(drop=True)
                            
                            simpan_stok()
                            simpan_transaksi()
                            
                            st.session_state.notif_tab2 = "🗑️ Transaksi dibatalkan. Stok Master telah dikembalikan seperti semula."
                            if "sel_edit_nama_tx" in st.session_state: del st.session_state["sel_edit_nama_tx"]
                            st.rerun()

        st.markdown("<br>", unsafe_allow_html=True)

        df_filter = st.session_state.df_transaksi.copy()
        df_filter["Tanggal_Str"] = df_filter["Waktu"].apply(lambda x: str(x).split()[0])
        df_filter["Tanggal_Obj"] = pd.to_datetime(df_filter["Tanggal_Str"]).dt.date
        
        df_filter = df_filter[(df_filter["Tanggal_Obj"] >= tgl_mulai) & (df_filter["Tanggal_Obj"] <= tgl_akhir)]
        df_filter = df_filter.sort_values(by="Waktu", ascending=False)
        
        if filter_barang != "Semua Barang":
            df_filter = df_filter[df_filter["Nama Barang"] == filter_barang]
            
        if df_filter.empty:
            st.info(f"Tidak ada aktivitas transaksi dari {tgl_mulai.strftime('%d %B %Y')} s/d {tgl_akhir.strftime('%d %B %Y')}.")
        else:
            df_display = df_filter[["Tanggal_Str", "Nama Barang", "Jenis", "ID Barang", "Jml Transaksi", "Pengambil", "Barang Tersedia"]]
            df_display = df_display.rename(columns={"Tanggal_Str": "Tanggal"})
            
            df_export = df_display.copy()
            df_export["Jenis"] = df_export["Jenis"].replace({'<span[^>]*>': '', '</span>': ''}, regex=True)

            jml_masuk = len(df_export[df_export["Jenis"] == "Masuk"])
            jml_keluar = len(df_export[df_export["Jenis"] == "Keluar"])
            teks_periode = f"{tgl_mulai.strftime('%d %b %Y')} - {tgl_akhir.strftime('%d %b %Y')}"

            col_d1, col_d2 = st.columns(2)

            output_excel = io.BytesIO()
            with pd.ExcelWriter(output_excel, engine='xlsxwriter') as writer:
                df_export.to_excel(writer, index=False, sheet_name="Laporan")
                worksheet = writer.sheets['Laporan']
                worksheet.set_column('A:A', 14)
                worksheet.set_column('B:B', 25)
                worksheet.set_column('C:C', 10)
                worksheet.set_column('D:D', 12)
                worksheet.set_column('E:E', 15)
                worksheet.set_column('F:F', 20)
                worksheet.set_column('G:G', 20)
                worksheet.autofilter('A1:F1')
            
            with col_d1:
                st.download_button(
                    label="📥 Download Laporan Excel", data=output_excel.getvalue(),
                    file_name=f"Laporan_Transaksi_{tgl_mulai}_sd_{tgl_akhir}.xlsx", mime="application/vnd.ms-excel",
                    use_container_width=True, key="btn_dl_excel"
                )

            doc = Document()
            table_hdr = doc.add_table(rows=1, cols=2)
            table_hdr.columns[0].width = Inches(1.0)
            table_hdr.columns[1].width = Inches(5.5)
            
            cell_logo = table_hdr.cell(0, 0)
            cell_text = table_hdr.cell(0, 1)
            
            for cell in [cell_logo, cell_text]:
                tcPr = cell._element.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), '12715B')
                tcPr.append(shd)
                
            if Path(nama_file_logo).is_file():
                p_logo = cell_logo.paragraphs[0]
                p_logo.alignment = 1 
                r_logo = p_logo.add_run()
                r_logo.add_picture(nama_file_logo, width=Inches(0.7))
                
            p_text = cell_text.paragraphs[0]
            run_title = p_text.add_run("SISTEM STOK BARANG\n")
            run_title.bold = True
            run_title.font.size = Pt(16)
            run_title.font.color.rgb = RGBColor(255, 255, 255)
            
            run_sub = p_text.add_run("REKTORAT UNIVERSITAS WINAYA MUKTI")
            run_sub.font.size = Pt(10)
            run_sub.font.color.rgb = RGBColor(255, 255, 255)
            
            doc.add_paragraph() 
            p_info = doc.add_paragraph()
            p_info.add_run("RINGKASAN TRANSAKSI\n").bold = True
            p_info.add_run(f"Periode Laporan : {teks_periode}\n")
            p_info.add_run(f"Total Aktivitas   : {len(df_export)} Transaksi ({jml_masuk} Masuk | {jml_keluar} Keluar)")
            doc.add_paragraph() 
            
            table_data = doc.add_table(rows=1, cols=len(df_export.columns))
            table_data.style = 'Table Grid'
            hdr_cells = table_data.rows[0].cells
            
            for i, col_name in enumerate(df_export.columns):
                hdr_cells[i].text = col_name
                hdr_cells[i].paragraphs[0].runs[0].bold = True
            
            for _, row in df_export.iterrows():
                row_cells = table_data.add_row().cells
                for i, value in enumerate(row): row_cells[i].text = str(value)
            
            output_word = io.BytesIO()
            doc.save(output_word)
            
            with col_d2:
                st.download_button(
                    label="📄 Download Laporan Word", data=output_word.getvalue(),
                    file_name=f"Laporan_Transaksi_{tgl_mulai}_sd_{tgl_akhir}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True, key="btn_dl_word"
                )

            st.markdown("<br>", unsafe_allow_html=True)
            st.markdown(df_display.to_html(classes="table-gudang", index=False, escape=False), unsafe_allow_html=True)
